"""
docreader.py —— 从各种格式的文件里，把纯文本抠出来。

RAG 的第一步是"读文档"。但真实的文档格式五花八门：
记事本存的 txt、写笔记用的 md、别人发来的 pdf、公司文件是 docx……
而不管前面是什么格式，**后面切块、算向量要的都只是一段纯文本**。

所以这个模块只干一件事：
    给我一个文件，我想办法把里面的字抠出来。

单独拆成一个文件，是因为"解析各种文档格式"和"跟模型对话"是两件不相干的事 ——
以后要加 pptx、html、epub，都在这里加，不用碰别的地方。
"""

import os

# 纯文本类：直接按文本读就行
TEXT_EXTENSIONS = (".txt", ".md")

# 需要额外库才能解析的
PDF_EXTENSIONS = (".pdf",)
DOCX_EXTENSIONS = (".docx",)
XLSX_EXTENSIONS = (".xlsx",)

# 对外暴露的"我支持哪些格式"
SUPPORTED_EXTENSIONS = (TEXT_EXTENSIONS + PDF_EXTENSIONS
                        + DOCX_EXTENSIONS + XLSX_EXTENSIONS)


def read_document(path):
    """读一个文件，把里面的文字抠出来。

    成功 → (文件名, 正文)
    失败 → (None, 一句人能看懂的错误说明)

    统一返回这个格式，是因为调用方（建库脚本、/file 命令）不该关心
    "这个文件到底是 pdf 还是 txt" —— 它们只要文字。
    """
    extension = os.path.splitext(path)[1].lower()

    if extension not in SUPPORTED_EXTENSIONS:
        supported = "、".join(SUPPORTED_EXTENSIONS)
        return None, f"不支持 {extension or '这种没有后缀名的文件'}，目前支持：{supported}"

    if os.path.isdir(path):
        return None, "这是个目录，不是文件"
    if not os.path.exists(path):
        return None, "找不到这个文件"

    name = os.path.basename(path)

    # 按格式分派到对应的解析函数
    if extension in PDF_EXTENSIONS:
        return _read_pdf(path, name)
    if extension in DOCX_EXTENSIONS:
        return _read_docx(path, name)
    if extension in XLSX_EXTENSIONS:
        return _read_xlsx(path, name)
    return _read_text(path, name)


def _read_text(path, name):
    """读纯文本文件。先按 UTF-8 试，不行再按 GBK 试。"""
    try:
        with open(path, "r", encoding="utf-8") as f:
            return name, f.read()
    except UnicodeDecodeError:
        # 中文 Windows 上，记事本存的老文件常常是 GBK 编码。
        # 这就是 requirements.txt 那个坑的同一个根：编码对不上而已。
        try:
            with open(path, "r", encoding="gbk") as f:
                return name, f.read()
        except (OSError, UnicodeDecodeError) as error:
            return None, f"读不了（{error}）"
    except OSError as error:
        return None, f"读不了（{error}）"


def _read_pdf(path, name):
    """从 PDF 里抠文字。

    ⚠️ 只对"文字型 PDF"有效 —— 也就是能用鼠标选中文字的 PDF。

    扫描件（本质是一堆图片）抠出来是空的，那需要 OCR，是另一回事。
    遇到这种情况会明确告诉你，而不是默默返回一个空文档。
    """
    try:
        from pypdf import PdfReader
    except ImportError:
        return None, "读 PDF 需要 pypdf，先运行：pip install pypdf"

    try:
        reader = PdfReader(path)
        pages = []
        for page in reader.pages:
            text = page.extract_text() or ""
            if text.strip():
                pages.append(text)
    except Exception as error:
        # 这里故意接得很宽：PDF 是个很乱的格式，pypdf 能抛出各种奇怪的异常。
        # 与其让一个坏文件把整个建库流程搞崩，不如跳过它、继续处理剩下的。
        return None, f"这个 PDF 解析失败（{type(error).__name__}: {error}）"

    if not pages:
        return None, "这个 PDF 里抠不出文字。如果是扫描件（图片），需要 OCR 才能读"

    return name, "\n\n".join(pages)


def _read_docx(path, name):
    """从 Word 文件里抠文字。

    只支持 .docx，不支持老的 .doc —— .doc 是完全不同的二进制格式，
    要读它得用别的库。
    """
    try:
        import docx
    except ImportError:
        return None, "读 Word 需要 python-docx，先运行：pip install python-docx"

    try:
        document = docx.Document(path)

        blocks = [p.text for p in document.paragraphs if p.text.strip()]

        # 表格也要抠 —— 很多规章制度的关键数字就藏在表格里，
        # 只读段落的话会漏掉一整块信息。
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    # 用竖线拼成一行，保留"这些字属于同一行"的关系
                    blocks.append(" | ".join(cells))
    except Exception as error:
        return None, f"这个 Word 文件解析失败（{type(error).__name__}: {error}）"

    if not blocks:
        return None, "这个 Word 文件里没有文字"

    return name, "\n".join(blocks)


def _read_xlsx(path, name):
    """从 Excel 文件里抠文字。

    做法：每一行拼成一行文本，单元格之间用 | 隔开，前面带上工作表名字。

    **为什么要保留"同一行的关系"**：
        表格里的数字单独拿出来是没有意义的 —— "128" 是什么？
        必须知道它属于哪一行的哪一列。用 | 拼起来，
        模型才能看出"年度预算 | 128 万元"是一对。

    只支持 .xlsx（新版 Excel）。老的 .xls 是另一种二进制格式，
    openpyxl 读不了，要另找库。
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        return None, "读 Excel 需要 openpyxl，先运行：pip install openpyxl"

    try:
        # read_only=True      → 大文件不把整个工作簿读进内存
        # data_only=True      → 读公式的**计算结果**，而不是公式本身
        #                        （不然你抠出来的是 "=SUM(A1:A9)"，没用）
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        # 跟 PDF 一样故意接得很宽：Excel 文件被改坏的情况五花八门，
        # 别让一个坏文件把整个建库流程搞崩
        return None, f"这个 Excel 解析失败（{type(error).__name__}: {error}）"

    try:
        sheets = []
        for sheet in workbook.worksheets:
            lines = []
            for row in sheet.iter_rows(values_only=True):
                # 空单元格直接跳过。
                # 代价是列的位置会错开，换来的是行更短、噪声更少 ——
                # 检索时我们关心的是"提到了哪些词"，不是表格排版。
                cells = [str(value).strip() for value in row
                         if value is not None and str(value).strip()]
                if cells:
                    lines.append(" | ".join(cells))

            if lines:
                sheets.append(f"【工作表：{sheet.title}】\n" + "\n".join(lines))
    finally:
        # read_only 模式下必须显式关闭，否则文件句柄会一直占着
        workbook.close()

    if not sheets:
        return None, "这个 Excel 里没有文字（可能全是空的）"

    return name, "\n\n".join(sheets)
